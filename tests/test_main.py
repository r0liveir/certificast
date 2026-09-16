import subprocess
from pathlib import Path
from unittest.mock import call, patch

import pytest
from pptx import Presentation

import certificast
from certificast.main import _ensure_empty_or_create


def make_template(path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 1_000_000, 1_000_000).text = "__NAME__ | __EVENT__"
    deck.save(str(path))


def fake_libreoffice(
    command: list[str], **options: object
) -> subprocess.CompletedProcess[str]:
    assert options["timeout"] == 122
    inputs = [Path(argument) for argument in command if argument.endswith(".pptx")]
    assert [Presentation(str(path)).slides[0].shapes[0].text for path in inputs] == [
        "Ana/Silva | Conference",
        "Bruno | Workshop",
    ]
    for path in inputs:
        path.with_suffix(".pdf").write_bytes(b"%PDF-1.4\n")
    return subprocess.CompletedProcess(command, 0, "", "")


def test_generate(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    csv_file = tmp_path / "people.csv"
    output = tmp_path / "output"
    make_template(template)
    csv_file.write_text(
        "Full name,EVENT\nAna/Silva,\nBruno,Workshop\n", encoding="utf-8"
    )

    with (
        patch("certificast.main.shutil.which", return_value="libreoffice"),
        patch("certificast.main.subprocess.run", side_effect=fake_libreoffice),
    ):
        result = certificast.generate(
            str(template),
            str(csv_file),
            str(output),
            {"NAME": "Full name"},
            {"EVENT": "Conference"},
            "{NAME}_{EVENT}",
        )

    assert result == [
        output / "0001-Ana_Silva_Conference.pdf",
        output / "0002-Bruno_Workshop.pdf",
    ]


def test_large_generation_uses_conversion_batches(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    output = tmp_path / "output"
    make_template(template)
    batch_sizes: list[int] = []

    def convert(
        command: list[str], **options: object
    ) -> subprocess.CompletedProcess[str]:
        inputs = [Path(argument) for argument in command if argument.endswith(".pptx")]
        batch_sizes.append(len(inputs))
        assert options["timeout"] == 120 + len(inputs)
        for path in inputs:
            path.with_suffix(".pdf").write_bytes(b"%PDF-1.4\n")
        return subprocess.CompletedProcess(command, 0, "", "")

    with (
        patch(
            "certificast.main.validate",
            return_value=[{"NAME": "Ana", "EVENT": "Conference"}] * 101,
        ),
        patch("certificast.main.shutil.which", return_value="libreoffice"),
        patch("certificast.main.subprocess.run", side_effect=convert),
    ):
        result = certificast.generate(str(template), "unused.csv", str(output))

    assert batch_sizes == [100, 1]
    assert len(result) == 101


def test_rejects_missing_value(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    csv_file = tmp_path / "people.csv"
    make_template(template)
    csv_file.write_text("NAME,EVENT\n,Workshop\n", encoding="utf-8")

    with pytest.raises(ValueError, match="missing value for 'NAME'"):
        certificast.validate(str(template), str(csv_file))


def test_rejects_bad_mapping_and_uncovered_variable(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    csv_file = tmp_path / "people.csv"
    make_template(template)
    csv_file.write_text("Full name\nAna\n", encoding="utf-8")

    with pytest.raises(ValueError, match="unknown variables"):
        certificast.validate(str(template), str(csv_file), {"WRONG": "Full name"})
    with pytest.raises(ValueError, match="missing required columns.*EVENT"):
        certificast.validate(str(template), str(csv_file), {"NAME": "Full name"})
    with pytest.raises(ValueError, match="unknown variables.*MISSING"):
        certificast.validate(
            str(template),
            str(csv_file),
            {"NAME": "Full name"},
            {"EVENT": "Event"},
            "{MISSING}",
        )


def test_rejects_split_placeholder(tmp_path: Path) -> None:
    template = tmp_path / "template.pptx"
    csv_file = tmp_path / "people.csv"
    make_template(template)
    deck = Presentation(str(template))
    paragraph = deck.slides[0].shapes[0].text_frame.paragraphs[0]
    paragraph.text = "__NA"
    paragraph.add_run().text = "ME__ | __EVENT__"
    deck.save(str(template))
    csv_file.write_text("NAME,EVENT\nAna,Conference\n", encoding="utf-8")

    with pytest.raises(ValueError, match="Placeholder __NAME__ is split"):
        certificast.validate(str(template), str(csv_file))


def test_output_must_be_empty(tmp_path: Path) -> None:
    output = tmp_path / "output"
    assert _ensure_empty_or_create(output) == output
    (output / "existing.pdf").touch()
    with pytest.raises(FileExistsError):
        _ensure_empty_or_create(output)


def test_pipeline_defers_validates_and_runs_jobs() -> None:
    pipeline = certificast.Pipeline()
    with (
        patch("certificast.main.validate") as validate,
        patch(
            "certificast.main.generate",
            side_effect=[[Path("first/0001.pdf")], [Path("second/0001.pdf")]],
        ) as generate,
    ):
        pipeline.add(
            input_template="one.pptx", input_file="one.csv", output_dir="first"
        )
        pipeline.add(
            input_template="two.pptx", input_file="two.csv", output_dir="second"
        )
        validate.assert_not_called()
        generate.assert_not_called()

        assert pipeline.run() == [Path("first/0001.pdf"), Path("second/0001.pdf")]

    assert validate.call_args_list == [
        call("one.pptx", "one.csv", None, None, None),
        call("two.pptx", "two.csv", None, None, None),
    ]
    assert generate.call_count == 2
