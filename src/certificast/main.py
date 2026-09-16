"""Minimal CSV-to-PDF generation, separate from the existing public API."""

import csv
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Not all patterns work in pptx, so we default to __VARIABLE__
VARIABLE_PATTERN = re.compile(r"__([A-Z0-9_]+)__")
type Job = tuple[
    str,
    str,
    str,
    dict[str, str] | None,
    dict[str, str] | None,
]


### --------------
# Helpers
### --------------
def _ensure_empty_or_create(dir_path: str | Path) -> Path:
    path = Path(dir_path)
    if not path.exists():
        path.mkdir(parents=True, exist_ok=True)
    elif not path.is_dir():
        raise NotADirectoryError(f"{path} exists but is not a directory.")
    elif any(path.iterdir()):
        raise FileExistsError(f"{path} exists and is not empty.")

    return path


def validate(
    input_template: str,
    input_file: str,
    columns_mapping: dict[str, str] | None = None,
    defaults: dict[str, str] | None = None,
) -> list[dict[str, str]]:
    """Validate inputs and return one resolved context per CSV row."""
    columns_mapping = columns_mapping or {}
    defaults = defaults or {}
    deck = Presentation(input_template)
    if len(deck.slides) != 1:
        raise ValueError("Template must contain exactly one slide.")
    if any(
        shape.has_table or shape.shape_type == MSO_SHAPE_TYPE.GROUP
        for shape in deck.slides[0].shapes
    ):
        raise ValueError("Tables and groups are not supported in this slice.")

    variables = {
        name
        for shape in deck.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for name in VARIABLE_PATTERN.findall(paragraph.text)
    }
    unknown_mappings = columns_mapping.keys() - variables
    if unknown_mappings:
        raise ValueError(
            f"Mappings contain unknown variables: {sorted(unknown_mappings)}"
        )

    contexts: list[dict[str, str]] = []
    with open(input_file, encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        headers = reader.fieldnames
        if not headers or len(set(headers)) != len(headers):
            raise ValueError("CSV requires unique column headers.")
        missing_columns = {
            columns_mapping.get(name, name)
            for name in variables
            if columns_mapping.get(name, name) not in headers and not defaults.get(name)
        }
        if missing_columns:
            raise ValueError(
                f"CSV is missing required columns: {sorted(missing_columns)}"
            )

        for index, row in enumerate(reader, start=1):
            if None in row or None in row.values():
                raise ValueError(f"CSV row {index} does not match its header.")
            values: dict[str, str] = {}
            for name in sorted(variables):
                value = row.get(columns_mapping.get(name, name)) or defaults.get(name)
                if value is None or not value.strip():
                    raise ValueError(f"Row {index}: missing value for {name!r}.")
                values[name] = value
            contexts.append(values)

    if not contexts:
        raise ValueError("Empty CSV given.")
    return contexts


def generate(
    input_template: str,
    input_file: str,
    output_dir: str,
    columns_mapping: dict[str, str] | None = None,
    defaults: dict[str, str] | None = None,
) -> list[Path]:
    """Generate numbered PDFs from a one-slide PPTX and UTF-8 CSV."""

    contexts = validate(input_template, input_file, columns_mapping, defaults)
    converter = shutil.which("libreoffice") or shutil.which("soffice")
    if converter is None:
        raise RuntimeError("Install LibreOffice and add it to PATH.")

    deck = Presentation(str(input_template))
    output_path = _ensure_empty_or_create(output_dir)

    # Extract all paragraph
    paragraphs = [
        (paragraph, paragraph.text)
        for shape in deck.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
    ]
    runs = [(run, run.text) for paragraph, _ in paragraphs for run in paragraph.runs]

    certificates: list[Path] = []

    # Take row data, render individual PPTX, invoke headless libreoff
    with tempfile.TemporaryDirectory(prefix="certificast-") as tmp_str:
        work_dir = Path(tmp_str)
        pptx_paths: list[Path] = []

        # Generate filled .pptx files
        for index, row in enumerate(contexts, start=1):
            # Always start from the template text, not the preceding row.
            for run, original_text in runs:
                run.text = original_text
                for key, val in row.items():
                    run.text = run.text.replace(f"__{key}__", val)

            out_pptx = work_dir / f"{index:04d}.pptx"
            deck.save(str(out_pptx))
            pptx_paths.append(out_pptx)

        # Convert all PPTX files to PDF in a single batch
        profile_uri = (work_dir / "profile").as_uri()
        cmd = [
            converter,
            f"-env:UserInstallation={profile_uri}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(work_dir),
            *[str(p) for p in pptx_paths],
        ]
        completed = subprocess.run(
            cmd, capture_output=True, text=True, timeout=120, check=True
        )

        # Validate and copy results to output_dir
        for pptx_path in pptx_paths:
            pdf_path = pptx_path.with_suffix(".pdf")

            # Check existence
            if not pdf_path.is_file():
                raise RuntimeError(
                    f"Conversion failed for {pdf_path.name}: {completed.stderr}"
                )

            # Check PDF magic bytes (%PDF-)
            with pdf_path.open("rb") as f:
                if f.read(5) != b"%PDF-":
                    raise RuntimeError(f"Corrupt PDF generated: {pdf_path.name}")

            # Copy to final destination
            target_path = output_path.resolve() / pdf_path.name
            shutil.copy2(pdf_path, target_path)
            certificates.append(target_path)

        return certificates


class Pipeline:
    """A collection of generation jobs executed later."""

    def __init__(self) -> None:
        self._jobs: list[Job] = []

    def add(
        self,
        *,
        input_template: str,
        input_file: str,
        output_dir: str,
        columns_mapping: dict[str, str] | None = None,
        defaults: dict[str, str] | None = None,
    ) -> None:
        self._jobs.append(
            (input_template, input_file, output_dir, columns_mapping, defaults)
        )

    def validate(self) -> None:
        for template, input_file, _, columns, defaults in self._jobs:
            validate(template, input_file, columns, defaults)

    def run(self) -> list[Path]:
        self.validate()
        return [certificate for job in self._jobs for certificate in generate(*job)]
